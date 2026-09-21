from __future__ import annotations

from pathlib import Path
from typing import Iterable

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .models import StudentRecord


HEADERS = [
    "座位号",
    "学号",
    "姓名",
    "校验备注",
]


def write_excel(records: Iterable[StudentRecord], output_path: Path) -> None:
    all_records = list(records)
    workbook = Workbook()
    normal_sheet = workbook.active
    normal_sheet.title = "学生信息"
    review_sheet = workbook.create_sheet("待复核")

    _write_sheet(normal_sheet, HEADERS, [record.as_row() for record in all_records])
    review_rows = [record.as_row() for record in all_records if record.status != "正常"]
    _write_sheet(review_sheet, HEADERS, review_rows)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)


def _write_sheet(sheet, headers: list[str], rows: list[list[object]]) -> None:
    sheet.append(headers)
    for row in rows:
        sheet.append(row)

    header_fill = PatternFill("solid", fgColor="1F4E78")
    for cell in sheet[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")

    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for column_cells in sheet.columns:
        width = min(
            60,
            max(12, max(len(str(cell.value or "")) for cell in column_cells) + 2),
        )
        sheet.column_dimensions[get_column_letter(column_cells[0].column)].width = width

    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
